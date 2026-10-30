"""Reporte mensual en Excel (app/services/reportes.py): reconstruccion de las
ocasiones de entrega a partir de la bitacora y armado del libro.

La base va simulada; el libro generado se vuelve a abrir con openpyxl para
afirmar sobre su contenido real (hojas, columnas, valores)."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.services import reportes
from app.services.reportes import (
    _devoluciones_por_entrega,
    _motivo_legible,
    _no_entregados_por_entrega,
    _ocasiones_por_entrega,
    _orden_tipo,
    generar_reporte_mensual_xlsx,
)
from tests.helpers import encabezado_bearer, fila_usuario_dashboard

UTC = timezone.utc


def dt(dia: int, mes: int = 1, anio: int = 2026, hora: int = 12) -> datetime:
    return datetime(anio, mes, dia, hora, 0, tzinfo=UTC)


def log(entrega: str, items: list[tuple[str, int]], cuando: datetime) -> dict:
    return {
        "entidad_id": entrega,
        "detalle": {"items": [{"id": i, "cantidad_entregada": c} for i, c in items]},
        "timestamp": cuando,
    }


# --------------------------------------------------------------------------
# Ocasiones de entrega reconstruidas desde los logs
# --------------------------------------------------------------------------


def test_la_primera_vez_que_aparece_un_item_cuenta_completa():
    ocasiones = _ocasiones_por_entrega([log("e1", [("a", 5)], dt(1))])
    assert ocasiones == {"e1": [(dt(1), 5)]}


def test_cada_ocasion_es_el_aumento_respecto_al_ultimo_valor_conocido():
    logs = [log("e1", [("a", 5)], dt(1)), log("e1", [("a", 8)], dt(2)), log("e1", [("a", 8), ("b", 2)], dt(3))]
    assert _ocasiones_por_entrega(logs) == {"e1": [(dt(1), 5), (dt(2), 3), (dt(3), 2)]}


def test_una_ocasion_suma_los_productos_del_mismo_guardado():
    assert _ocasiones_por_entrega([log("e1", [("a", 4), ("b", 6)], dt(1))]) == {"e1": [(dt(1), 10)]}


@pytest.mark.parametrize(
    "secuencia",
    [
        [5, 5],  # guardado que solo cambio una nota: delta 0
        [5, 3],  # una devolucion: delta negativo
    ],
)
def test_un_delta_cero_o_negativo_no_es_ocasion_de_entrega(secuencia):
    logs = [log("e1", [("a", v)], dt(n + 1)) for n, v in enumerate(secuencia)]
    assert _ocasiones_por_entrega(logs) == {"e1": [(dt(1), 5)]}


def test_tras_una_devolucion_la_siguiente_entrega_cuenta_contra_el_valor_bajo():
    logs = [log("e1", [("a", 5)], dt(1)), log("e1", [("a", 3)], dt(2)), log("e1", [("a", 7)], dt(3))]
    assert _ocasiones_por_entrega(logs) == {"e1": [(dt(1), 5), (dt(3), 4)]}


def test_las_entregas_no_se_mezclan_entre_si():
    logs = [log("e1", [("a", 5)], dt(1)), log("e2", [("a", 9)], dt(1))]
    assert _ocasiones_por_entrega(logs) == {"e1": [(dt(1), 5)], "e2": [(dt(1), 9)]}


@pytest.mark.parametrize(
    "detalle",
    [None, {}, {"items": None}, {"items": []}, {"items": [{"id": None, "cantidad_entregada": 3}]}, {"items": [{"id": "a"}]}],
)
def test_logs_sin_items_utiles_se_ignoran(detalle):
    assert _ocasiones_por_entrega([{"entidad_id": "e1", "detalle": detalle, "timestamp": dt(1)}]) == {}


# --------------------------------------------------------------------------
# Textos legibles
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "motivo, detalle, esperado",
    [
        ("danado", None, "Dañado"),
        ("otro", "se mojo en el camino", "se mojo en el camino"),  # el texto del operador gana
        ("otro", None, "Otro"),
        ("otro", "", "Otro"),
        ("danado", "ignorado", "Dañado"),  # el detalle solo aplica a "otro"
        ("motivo_viejo", None, "motivo_viejo"),  # valor desconocido: se muestra tal cual
    ],
)
def test_motivo_legible(motivo, detalle, esperado):
    assert _motivo_legible(reportes._MOTIVOS_LEGIBLES, motivo, detalle) == esperado


def test_cada_motivo_y_resolucion_del_modelo_tiene_texto_legible():
    from app.models.devolucion import MotivoDevolucion, ResolucionDevolucion
    from app.models.no_entregado import MotivoNoEntregado

    assert {m.value for m in MotivoDevolucion} <= set(reportes._MOTIVOS_LEGIBLES)
    assert {r.value for r in ResolucionDevolucion} <= set(reportes._RESOLUCIONES_LEGIBLES)
    assert {m.value for m in MotivoNoEntregado} <= set(reportes._MOTIVOS_NO_ENTREGADO_LEGIBLES)


def test_resumen_de_devoluciones_por_entrega():
    filas = [
        {"entrega_id": "e1", "creado_at": dt(5), "cantidad": 2, "motivo": "danado", "motivo_detalle": None, "resolucion": "reposicion"},
        {"entrega_id": "e1", "creado_at": dt(9), "cantidad": 1, "motivo": "otro", "motivo_detalle": "no gusto", "resolucion": "reembolso"},
        {"entrega_id": "e2", "creado_at": dt(6), "cantidad": 4, "motivo": "vencido", "motivo_detalle": None, "resolucion": "no_lo_lleva"},
    ]
    assert _devoluciones_por_entrega(filas) == {
        "e1": ["05/01/2026: 2 und (Dañado, Reposición)", "09/01/2026: 1 und (no gusto, Reembolso)"],
        "e2": ["06/01/2026: 4 und (Vencido, No lo lleva)"],
    }


def test_resumen_de_no_entregados_por_entrega():
    filas = [
        {"entrega_id": "e1", "creado_at": dt(5), "cantidad": 3, "motivo": "facturado_de_mas", "motivo_detalle": None},
        {"entrega_id": "e1", "creado_at": dt(6), "cantidad": 1, "motivo": "sin_existencia", "motivo_detalle": None},
    ]
    assert _no_entregados_por_entrega(filas) == {
        "e1": ["05/01/2026: 3 und (Facturado de más)", "06/01/2026: 1 und (Sin existencia)"]
    }


def test_orden_de_tipos_conocidos_primero_y_luego_alfabetico():
    tipos = ["ZZZ", "RSF", "AAA", "FEI", "RM2", "TB9", "RM3"]
    assert sorted(tipos, key=_orden_tipo) == ["FEI", "TB9", "RM3", "RM2", "RSF", "AAA", "ZZZ"]


# --------------------------------------------------------------------------
# Libro completo
# --------------------------------------------------------------------------


def entrega(id: str, cuando: datetime, tipo: str = "FEI", numero: str = "100", sede: str = "s1") -> dict:
    return {"id": id, "capturado_at": cuando, "tipo": tipo, "indicativo_numero": numero, "sede_origen_id": sede}


def guionar(conn, entregas, cantidades=(), logs=(), devoluciones=(), no_entregados=()) -> None:
    conn.responder("fetch", "from entregas order by", list(entregas))
    conn.responder("fetch", "from entrega_items group by", [{"entrega_id": e, "cantidad_total": c} for e, c in cantidades])
    conn.responder("fetch", "from logs", list(logs))
    conn.responder("fetch", "from devoluciones", list(devoluciones))
    conn.responder("fetch", "from items_no_entregados", list(no_entregados))


async def libro(**kwargs):
    return load_workbook(io.BytesIO(await generar_reporte_mensual_xlsx(**kwargs)))


def fila(ws, n: int, columnas: int = 8, desde: int = 1) -> list:
    return [ws.cell(row=n, column=desde + i).value for i in range(columnas)]


async def test_una_hoja_por_mes_en_orden_cronologico_con_nombre_en_espanol(pool, conn):
    guionar(
        conn,
        [entrega("e3", dt(3, 3)), entrega("e1", dt(10, 1, 2025), numero="1"), entrega("e2", dt(4, 2), numero="2")],
    )
    assert (await libro()).sheetnames == ["Enero 2025", "Febrero 2026", "Marzo 2026"]


async def test_bloque_por_tipo_con_encabezados_y_valores(pool, conn):
    guionar(
        conn,
        [entrega("e1", dt(5), "FEI", "100")],
        cantidades=[("e1", 12)],
        logs=[log("e1", [("a", 5)], dt(6)), log("e1", [("a", 12)], dt(9))],
        devoluciones=[
            {"entrega_id": "e1", "creado_at": dt(10), "cantidad": 2, "motivo": "danado", "motivo_detalle": None, "resolucion": "reposicion"}
        ],
        no_entregados=[
            {"entrega_id": "e1", "creado_at": dt(11), "cantidad": 1, "motivo": "otro", "motivo_detalle": "sobraba"}
        ],
    )
    ws = (await libro())["Enero 2026"]
    assert ws["A1"].value == "FEI"
    assert fila(ws, 2) == [
        "Fecha", "Número", "Cantidad entregada", "N° entregas",
        "Fechas entregas", "Cantidades entregas", "Devoluciones", "No entregados",
    ]
    assert fila(ws, 3) == [
        "05/01/2026", "100", 12, 2, "06/01/2026, 09/01/2026", "5, 7",
        "10/01/2026: 2 und (Dañado, Reposición)", "11/01/2026: 1 und (sobraba)",
    ]


async def test_varios_tipos_van_lado_a_lado_con_una_columna_en_blanco_y_en_orden(pool, conn):
    guionar(
        conn,
        [entrega("e1", dt(5), "ZZZ", "1"), entrega("e2", dt(5), "FEI", "2"), entrega("e3", dt(5), "RM3", "3")],
    )
    ws = (await libro())["Enero 2026"]
    # FEI en A..H, (I en blanco), RM3 en J..Q, (R en blanco), ZZZ en S..Z
    assert (ws["A1"].value, ws["J1"].value, ws["S1"].value) == ("FEI", "RM3", "ZZZ")
    assert ws["I1"].value is None and ws["I3"].value is None
    assert (ws["B3"].value, ws["K3"].value, ws["T3"].value) == ("2", "3", "1")


async def test_las_entregas_de_un_tipo_van_ordenadas_por_fecha(pool, conn):
    guionar(conn, [entrega("e1", dt(9), numero="9"), entrega("e2", dt(2), numero="2"), entrega("e3", dt(5), numero="5")])
    ws = (await libro())["Enero 2026"]
    assert [ws[f"B{n}"].value for n in (3, 4, 5)] == ["2", "5", "9"]


async def test_tipo_vacio_se_agrupa_como_sin_tipo_y_numero_nulo_queda_vacio(pool, conn):
    guionar(conn, [entrega("e1", dt(5), "  ", "")])
    ws = (await libro())["Enero 2026"]
    assert ws["A1"].value == "SIN TIPO"
    # openpyxl no distingue "" de "sin valor": ambos se leen como None.
    assert ws["B3"].value in (None, "")


async def test_cantidad_sin_items_es_cero_y_sin_ocasiones_las_columnas_quedan_vacias(pool, conn):
    guionar(conn, [entrega("e1", dt(5))])
    ws = (await libro())["Enero 2026"]
    assert ws["C3"].value == 0 and ws["D3"].value == 0
    assert ws["E3"].value in (None, "") and ws["F3"].value in (None, "")


async def test_el_filtro_de_sede_deja_solo_sus_entregas_y_sus_registros(pool, conn):
    guionar(
        conn,
        [entrega("e1", dt(5), numero="1", sede="s1"), entrega("e2", dt(5), numero="2", sede="s2")],
        cantidades=[("e1", 3), ("e2", 99)],
        logs=[log("e1", [("a", 3)], dt(6)), log("e2", [("a", 99)], dt(6))],
    )
    ws = (await libro(sede_id="s1"))["Enero 2026"]
    assert ws.max_row == 3
    assert (ws["B3"].value, ws["C3"].value, ws["D3"].value) == ("1", 3, 1)


async def test_registros_de_entregas_ajenas_al_filtro_no_se_cuelan(pool, conn):
    guionar(
        conn,
        [entrega("e1", dt(5), sede="s1"), entrega("e2", dt(5), numero="2", sede="s2")],
        devoluciones=[
            {"entrega_id": "e2", "creado_at": dt(7), "cantidad": 1, "motivo": "danado", "motivo_detalle": None, "resolucion": "reembolso"}
        ],
        no_entregados=[{"entrega_id": "e2", "creado_at": dt(7), "cantidad": 1, "motivo": "otro", "motivo_detalle": "x"}],
    )
    ws = (await libro(sede_id="s1"))["Enero 2026"]
    assert ws["G3"].value in (None, "") and ws["H3"].value in (None, "")


async def test_el_nombre_de_hoja_nunca_pasa_de_31_caracteres(pool, conn):
    guionar(conn, [entrega("e1", dt(1, 9))])
    (nombre,) = (await libro()).sheetnames
    assert nombre == "Septiembre 2026" and len(nombre) <= 31


async def test_el_reporte_solo_lee_nunca_escribe(pool, conn):
    guionar(conn, [entrega("e1", dt(5))])
    await generar_reporte_mensual_xlsx()
    assert all(c[0] == "fetch" for c in conn.llamadas)
    assert conn.inserts_log() == []


async def test_el_reporte_mensual_exige_una_sesion_del_dashboard(cliente, pool, conn):
    assert (await cliente.get("/entregas/export.xlsx")).status_code == 401
    assert conn.llamadas == []


@pytest.mark.parametrize("rol", ["admin", "supervisor", "consulta"])
async def test_el_reporte_mensual_lo_descargan_todos_los_roles_con_cabeceras_de_excel(cliente, pool, conn, rol):
    f = fila_usuario_dashboard(rol=rol)
    conn.responder("fetchrow", "from usuarios_dashboard where id", f)
    guionar(conn, [entrega("e1", dt(5))])
    res = await cliente.get("/entregas/export.xlsx", headers=encabezado_bearer(f))
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "reporte_mensual.xlsx" in res.headers["content-disposition"]
    assert load_workbook(io.BytesIO(res.content)).sheetnames == ["Enero 2026"]


async def test_el_reporte_mensual_pasa_el_filtro_de_sede(cliente, pool, conn):
    f = fila_usuario_dashboard(rol="consulta")
    conn.responder("fetchrow", "from usuarios_dashboard where id", f)
    guionar(conn, [entrega("e1", dt(5), sede="s1"), entrega("e2", dt(5), numero="2", sede="s2")])
    res = await cliente.get("/entregas/export.xlsx", params={"sede_id": "s2"}, headers=encabezado_bearer(f))
    ws = load_workbook(io.BytesIO(res.content))["Enero 2026"]
    assert ws.max_row == 3 and ws["B3"].value == "2"
