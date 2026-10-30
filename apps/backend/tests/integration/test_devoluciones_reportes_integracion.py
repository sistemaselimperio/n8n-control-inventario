"""Devoluciones, productos no entregados, reporte mensual y ranking contra
Postgres real (la IA sigue simulada)."""

from __future__ import annotations

import asyncio
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from tests import respuestas_vision as r
from tests.helpers import payload_entrega
from tests.integration.conftest import crear_entrega
from tests.integration.datos_admin import confirmar_item, entrega_con_item, sesion_dashboard

pytestmark = pytest.mark.integration


def cuerpo_devolucion(item_id: str, cantidad: int = 2, resolucion: str = "reposicion", **extra) -> dict:
    cuerpo = {
        "item_id": item_id, "cantidad": cantidad, "motivo": "danado", "resolucion": resolucion,
        "operador_id": "op-1", "sede_id": "sede-x",
    }
    cuerpo.update(extra)
    return cuerpo


def cuerpo_no_entregado(item_id: str, cantidad: int = 2, **extra) -> dict:
    cuerpo = {"item_id": item_id, "cantidad": cantidad, "motivo": "facturado_de_mas", "operador_id": "op-1", "sede_id": "sede-x"}
    cuerpo.update(extra)
    return cuerpo


async def estado_item(pg, item_id: str) -> tuple[int, int, int]:
    f = await pg.fetchrow(
        "select cantidad_entregada, cantidad_pendiente, cantidad_no_entregada from entrega_items where id = $1::uuid",
        item_id,
    )
    return f["cantidad_entregada"], f["cantidad_pendiente"], f["cantidad_no_entregada"]


# --------------------------------------------------------------------------
# Devoluciones
# --------------------------------------------------------------------------


async def test_reposicion_baja_lo_entregado_sube_lo_pendiente_y_deja_registro(cliente, pg):
    entrega_id, item_id = await entrega_con_item(pg)
    await confirmar_item(pg, item_id, entregada=5, pendiente=0)
    await pg.execute("update entrega_items set entregado_por = 'op-1', entregado_en_punto_venta = true where id = $1::uuid", item_id)

    res = await cliente.post(f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_id, 2, "reposicion"))
    assert res.status_code == 200 and res.json()["item"]["cantidad_entregada"] == 3
    assert await estado_item(pg, item_id) == (3, 2, 0)
    fila = await pg.fetchrow("select entregado_por, entregado_en_punto_venta from entrega_items where id = $1::uuid", item_id)
    assert (fila["entregado_por"], fila["entregado_en_punto_venta"]) == (None, False)  # el item se reabre
    dev = await pg.fetchrow("select * from devoluciones")
    assert (dev["cantidad"], dev["motivo"], dev["resolucion"], dev["operador_id"]) == (2, "danado", "reposicion", "op-1")
    assert [f["evento"] for f in await pg.fetch("select evento from logs")] == ["devolucion_registrada"]


@pytest.mark.parametrize("resolucion", ["reembolso", "no_lo_lleva"])
async def test_reembolso_y_no_lo_lleva_no_vuelven_a_pendiente(cliente, pg, resolucion):
    entrega_id, item_id = await entrega_con_item(pg)
    await confirmar_item(pg, item_id, entregada=5, pendiente=1)
    res = await cliente.post(f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_id, 2, resolucion))
    assert res.status_code == 200
    assert await estado_item(pg, item_id) == (3, 1, 0)


async def test_no_se_puede_devolver_mas_de_lo_entregado_y_nada_cambia(cliente, pg):
    entrega_id, item_id = await entrega_con_item(pg)
    await confirmar_item(pg, item_id, entregada=1, pendiente=0)
    res = await cliente.post(f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_id, 2))
    assert res.status_code == 422 and "solo tiene 1 entregado" in res.json()["detail"]
    assert await estado_item(pg, item_id) == (1, 0, 0)
    assert await pg.fetchval("select count(*) from devoluciones") == 0
    assert await pg.fetchval("select count(*) from logs") == 0


async def test_un_item_recien_escaneado_sin_confirmar_no_admite_devoluciones(cliente, pg):
    """crear_entrega deja actualizado_at = creado_at: nadie confirmo nada todavia."""
    entrega_id, (item_id,) = await crear_entrega(pg, numero="77", pendientes=(5,))
    await pg.execute("update entrega_items set cantidad_entregada = 5 where id = $1::uuid", item_id)
    res = await cliente.post(f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_id, 1))
    assert res.status_code == 422 and "todavia no fue confirmado" in res.json()["detail"]
    assert await estado_item(pg, item_id) == (5, 5, 0)


async def test_un_item_de_otra_entrega_no_se_puede_devolver(cliente, pg):
    entrega_a, _ = await entrega_con_item(pg)
    _, item_b = await entrega_con_item(pg)
    await confirmar_item(pg, item_b, entregada=5, pendiente=0)
    res = await cliente.post(f"/entregas/{entrega_a}/devoluciones", json=cuerpo_devolucion(item_b, 1))
    assert res.status_code == 422 and "no existe en esta entrega" in res.json()["detail"]
    assert await estado_item(pg, item_b) == (5, 0, 0)


async def test_dos_devoluciones_simultaneas_no_devuelven_mas_de_lo_entregado(cliente, pg):
    entrega_id, item_id = await entrega_con_item(pg)
    await confirmar_item(pg, item_id, entregada=4, pendiente=0)
    respuestas = await asyncio.gather(
        *[cliente.post(f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_id, 3)) for _ in range(2)]
    )
    assert sorted(r_.status_code for r_ in respuestas) == [200, 422]
    assert await estado_item(pg, item_id) == (1, 3, 0)
    assert await pg.fetchval("select count(*) from devoluciones") == 1


async def test_devolucion_de_entrega_inexistente_da_404(cliente, pg):
    inexistente = "99999999-9999-9999-9999-999999999999"
    res = await cliente.post(f"/entregas/{inexistente}/devoluciones", json=cuerpo_devolucion(inexistente))
    assert res.status_code == 404


# --------------------------------------------------------------------------
# Productos no entregados
# --------------------------------------------------------------------------


async def test_no_entregado_mueve_de_pendiente_sin_marcar_el_item_como_confirmado(cliente, pg):
    entrega_id, (item_id,) = await crear_entrega(pg, numero="88", pendientes=(5,))
    res = await cliente.post(f"/entregas/{entrega_id}/no-entregados", json=cuerpo_no_entregado(item_id, 2))
    assert res.status_code == 200
    assert res.json()["item"]["cantidad_pendiente"] == 3 and res.json()["item"]["confirmado"] is False
    assert await estado_item(pg, item_id) == (0, 3, 2)
    # No toca actualizado_at: la primera confirmacion de cantidades sigue funcionando.
    assert await pg.fetchval("select actualizado_at = creado_at from entrega_items where id = $1::uuid", item_id)
    fila = await pg.fetchrow("select * from items_no_entregados")
    assert (fila["cantidad"], fila["motivo"]) == (2, "facturado_de_mas")
    assert [f["evento"] for f in await pg.fetch("select evento from logs")] == ["item_no_entregado"]


async def test_no_entregado_mayor_a_lo_pendiente_da_422_sin_cambios(cliente, pg):
    entrega_id, (item_id,) = await crear_entrega(pg, numero="89", pendientes=(1,))
    res = await cliente.post(f"/entregas/{entrega_id}/no-entregados", json=cuerpo_no_entregado(item_id, 2))
    assert res.status_code == 422 and "solo tiene 1 pendiente" in res.json()["detail"]
    assert await estado_item(pg, item_id) == (0, 1, 0)


@pytest.mark.parametrize("rol, esperado", [("faia_viewer", 403), ("punto_venta", 403), ("operador", 200)])
async def test_no_entregado_segun_el_rol_del_empleado(cliente, pg, rol, esperado):
    entrega_id, (item_id,) = await crear_entrega(pg, numero="90", pendientes=(5,))
    emp = str(await pg.fetchval("insert into empleados (nombre, sede_id, rol) values ('X', 'sede-x', $1) returning id", rol))
    res = await cliente.post(f"/entregas/{entrega_id}/no-entregados", json=cuerpo_no_entregado(item_id, 1, operador_id=emp))
    assert res.status_code == esperado
    if esperado == 403:
        assert await estado_item(pg, item_id) == (0, 5, 0)


async def test_dos_no_entregados_simultaneos_no_superan_lo_pendiente(cliente, pg):
    entrega_id, (item_id,) = await crear_entrega(pg, numero="91", pendientes=(4,))
    respuestas = await asyncio.gather(
        *[cliente.post(f"/entregas/{entrega_id}/no-entregados", json=cuerpo_no_entregado(item_id, 3)) for _ in range(2)]
    )
    assert sorted(r_.status_code for r_ in respuestas) == [200, 422]
    assert await estado_item(pg, item_id) == (0, 1, 3)


# --------------------------------------------------------------------------
# Reporte mensual con un flujo real de principio a fin
# --------------------------------------------------------------------------


async def test_el_reporte_refleja_entregas_devoluciones_y_no_entregados_reales(cliente, pg, vision):
    headers = await sesion_dashboard(pg, "consulta1", "consulta")
    vision.responder({**r.VARIOS_ITEMS, "tipo": "XYZ", "indicativo_numero": "7001"})
    alta = await cliente.post("/entregas/procesar", json=payload_entrega(hash_evidencia="h-rep"))
    assert alta.status_code == 201, alta.text
    nueva = alta.json()
    entrega_id, items = nueva["id"], nueva["items"]
    item_a, item_b = items[0]["id"], items[1]["id"]
    contexto = {"operador_id": "op1", "sede_id": "sede-1"}

    # Paso 2 (confirmar cantidades) y dos visitas de entrega.
    await cliente.patch(
        f"/entregas/{entrega_id}/items",
        json={"items": [{"id": item_a, "cantidad_pendiente": 4}, {"id": item_b, "cantidad_pendiente": 10}], **contexto},
    )
    await cliente.patch(f"/entregas/{entrega_id}/items", json={"items": [{"id": item_a, "entregado_hoy": 3}], **contexto})
    await cliente.patch(
        f"/entregas/{entrega_id}/items",
        json={"items": [{"id": item_a, "entregado_hoy": 1}, {"id": item_b, "entregado_hoy": 2}], **contexto},
    )
    dev = await cliente.post(
        f"/entregas/{entrega_id}/devoluciones", json=cuerpo_devolucion(item_a, 1, "reembolso", sede_id="sede-1")
    )
    assert dev.status_code == 200, dev.text
    no = await cliente.post(f"/entregas/{entrega_id}/no-entregados", json=cuerpo_no_entregado(item_b, 1, sede_id="sede-1"))
    assert no.status_code == 200, no.text

    res = await cliente.get("/entregas/export.xlsx", headers=headers)
    assert res.status_code == 200
    (hoja,) = load_workbook(io.BytesIO(res.content)).worksheets
    assert hoja["A1"].value == "XYZ"
    # Entregado = (3 + 1) + 2 - 1 devuelta = 5, repartido en dos ocasiones reales.
    assert hoja["B3"].value == "7001"
    assert hoja["C3"].value == 5
    assert hoja["D3"].value == 2
    assert hoja["F3"].value == "3, 3"
    assert "1 und (Dañado, Reembolso)" in hoja["G3"].value
    assert "1 und (Facturado de más)" in hoja["H3"].value


async def test_el_reporte_con_la_base_vacia_responde_200_con_una_hoja_valida(cliente, pg):
    """Regresion: antes era HTTP 500 (libro sin hojas)."""
    headers = await sesion_dashboard(pg, "consulta1", "consulta")
    res = await cliente.get("/entregas/export.xlsx", headers=headers)
    assert res.status_code == 200
    assert load_workbook(io.BytesIO(res.content)).sheetnames == ["Sin datos"]


async def test_el_reporte_filtrado_por_una_sede_sin_entregas_tambien_responde_200(cliente, pg):
    headers = await sesion_dashboard(pg, "consulta1", "consulta")
    await entrega_con_item(pg, sede_origen_id="sede-a")
    res = await cliente.get("/entregas/export.xlsx", params={"sede_id": "sede-b"}, headers=headers)
    assert res.status_code == 200
    assert load_workbook(io.BytesIO(res.content)).sheetnames == ["Sin datos"]


async def test_el_reporte_separa_por_mes_y_por_sede(cliente, pg):
    headers = await sesion_dashboard(pg, "consulta1", "consulta")
    for fecha, sede in (("2026-01-10T12:00:00+00:00", "sede-a"), ("2026-02-10T12:00:00+00:00", "sede-a"), ("2026-02-11T12:00:00+00:00", "sede-b")):
        await entrega_con_item(pg, capturado_at=datetime.fromisoformat(fecha), sede_origen_id=sede)
    todo = load_workbook(io.BytesIO((await cliente.get("/entregas/export.xlsx", headers=headers)).content))
    assert todo.sheetnames == ["Enero 2026", "Febrero 2026"]
    solo_a = load_workbook(
        io.BytesIO((await cliente.get("/entregas/export.xlsx", params={"sede_id": "sede-a"}, headers=headers)).content)
    )
    assert solo_a["Febrero 2026"].max_row == 3  # encabezado, titulos y una sola entrega


# --------------------------------------------------------------------------
# Ranking de productos
# --------------------------------------------------------------------------


async def _item(pg, entrega_id: str, descripcion: str, entregada: int) -> None:
    await pg.execute(
        "insert into entrega_items (entrega_id, descripcion, cantidad_entregada) values ($1::uuid, $2, $3)",
        entrega_id, descripcion, entregada,
    )


async def _entrega_en(pg, fecha: str, sede: str = "sede-a", estado: str = "procesada") -> str:
    entrega_id, _ = await entrega_con_item(pg, capturado_at=datetime.fromisoformat(fecha), sede_origen_id=sede, estado=estado)
    await pg.execute("delete from entrega_items where entrega_id = $1::uuid", entrega_id)
    return entrega_id


RANGO = {"desde": "2026-01-01T00:00:00+00:00", "hasta": "2026-01-31T23:59:59+00:00"}


async def test_ranking_agrega_filtra_por_rango_sede_y_excluye_duplicados(cliente, pg):
    dentro_a = await _entrega_en(pg, "2026-01-10T12:00:00+00:00", "sede-a")
    dentro_b = await _entrega_en(pg, "2026-01-12T12:00:00+00:00", "sede-b")
    fuera = await _entrega_en(pg, "2026-02-12T12:00:00+00:00", "sede-a")
    duplicada = await _entrega_en(pg, "2026-01-15T12:00:00+00:00", "sede-a", estado="duplicado_bloqueado")
    for entrega, cantidad in ((dentro_a, 5), (dentro_b, 3), (fuera, 100), (duplicada, 100)):
        await _item(pg, entrega, "75936 SAL BLANCA", cantidad)
    await _item(pg, dentro_a, "AZUCAR", 1)
    await pg.execute("insert into productos (codigo, nombre) values ('75936', 'Sal fina 40 kg')")

    todo = (await cliente.get("/ranking/productos", params=RANGO)).json()
    mas = {i["codigo"] or i["nombre"]: i for i in todo["mas_vendidos"]}
    assert mas["75936"]["cantidad_total"] == 8 and mas["75936"]["entregas_count"] == 2
    assert mas["75936"]["nombre"] == "Sal fina 40 kg"  # el catalogo corrige el nombre
    assert todo["menos_vendidos"][0]["nombre"] == "AZUCAR"

    solo_a = (await cliente.get("/ranking/productos", params={**RANGO, "sede_id": "sede-a"})).json()
    assert {i["codigo"] or i["nombre"]: i["cantidad_total"] for i in solo_a["mas_vendidos"]} == {"75936": 5, "AZUCAR": 1}


async def test_ranking_el_limite_aplica_a_ambas_listas(cliente, pg):
    entrega = await _entrega_en(pg, "2026-01-10T12:00:00+00:00")
    for n in range(1, 6):
        await _item(pg, entrega, f"1000{n} PRODUCTO {n}", n)
    res = (await cliente.get("/ranking/productos", params={**RANGO, "limit": 2})).json()
    assert [i["cantidad_total"] for i in res["mas_vendidos"]] == [5, 4]
    assert [i["cantidad_total"] for i in res["menos_vendidos"]] == [1, 2]


async def test_ranking_sin_datos_devuelve_listas_vacias(cliente, pg):
    res = await cliente.get("/ranking/productos", params=RANGO)
    assert res.status_code == 200 and res.json()["mas_vendidos"] == [] and res.json()["menos_vendidos"] == []
